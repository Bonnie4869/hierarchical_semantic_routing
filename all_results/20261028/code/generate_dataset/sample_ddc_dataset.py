"""Create a balanced three-level DDC training workbook.

Sampling rules
--------------
* DL1: exactly 500 rows per existing DDC node. If the node has fewer than
  500 DL1 rows, fill the deficit from its immediate DL2 children.
* DL2: exactly 250 rows per existing DDC node. If the node has fewer than
  250 DL2 rows, fill the deficit from its immediate DL3 children.
* DL3: remove nodes with fewer than 10 rows; otherwise retain at most 50
  randomly selected rows per node.

Rows copied upward are relabelled to the target parent DDC. ``source_DL`` and
``source_DDC`` preserve their origin. Sampling is reproducible with --seed.
The script uses a temporary SQLite database and streaming XLSX I/O so the full
set of long descriptions does not need to remain in memory.
"""

import argparse
import random
import sqlite3
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TARGETS = {1: 500, 2: 250, 3: 50}
DL3_MINIMUM = 10
SHEET_COLORS = {1: "1F4E78", 2: "4472C4", 3: "5B9BD5"}
OUTPUT_HEADERS = (
    "DDC",
    "Title",
    "description",
    "remark",
    "source_DL",
    "source_DDC",
)


def parse_args():
    base_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=base_dir / "train_stratified_DL1_DL2_DL3.xlsx",
        help="Input workbook with DL1, DL2 and DL3 sheets",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=base_dir / "train_stratified_balanced.xlsx",
        help="Balanced output workbook",
    )
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow replacing an existing output workbook",
    )
    parser.add_argument(
        "--keep-temp",
        action="store_true",
        help="Keep the temporary SQLite database for inspection",
    )
    return parser.parse_args()


def normalize_three_digit_ddc(value, sheet_name, row_number):
    if value is None:
        raise ValueError(f"{sheet_name} row {row_number}: empty DDC")
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if not text.isdigit() or not 1 <= len(text) <= 3:
        raise ValueError(
            f"{sheet_name} row {row_number}: invalid DDC {value!r}"
        )
    number = int(text)
    if not 0 <= number <= 999:
        raise ValueError(
            f"{sheet_name} row {row_number}: DDC out of range {value!r}"
        )
    return f"{number:03d}"


def create_database(path):
    connection = sqlite3.connect(path)
    connection.execute("PRAGMA journal_mode=OFF")
    connection.execute("PRAGMA synchronous=OFF")
    connection.execute("PRAGMA temp_store=FILE")
    connection.execute(
        """
        CREATE TABLE rows (
            id INTEGER PRIMARY KEY,
            level INTEGER NOT NULL,
            ddc TEXT NOT NULL,
            parent_dl1 TEXT NOT NULL,
            parent_dl2 TEXT NOT NULL,
            title TEXT,
            description TEXT,
            remark TEXT,
            random_key REAL NOT NULL
        )
        """
    )
    return connection


def load_input_into_database(input_path, connection, seed):
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    missing = [name for name in ("DL1", "DL2", "DL3") if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Input workbook is missing sheets: {missing}")

    rng = random.Random(seed)
    counts = {1: 0, 2: 0, 3: 0}
    insert_sql = (
        "INSERT INTO rows "
        "(level,ddc,parent_dl1,parent_dl2,title,description,remark,random_key) "
        "VALUES (?,?,?,?,?,?,?,?)"
    )

    for level in (1, 2, 3):
        sheet_name = f"DL{level}"
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = tuple(next(rows))
        required = ("DDC", "Title", "description", "remark")
        missing_columns = [column for column in required if column not in headers]
        if missing_columns:
            raise ValueError(
                f"{sheet_name} is missing columns: {missing_columns}; found {headers}"
            )
        indexes = {column: headers.index(column) for column in required}

        batch = []
        for row_number, row in enumerate(rows, 2):
            ddc = normalize_three_digit_ddc(
                row[indexes["DDC"]], sheet_name, row_number
            )
            batch.append(
                (
                    level,
                    ddc,
                    ddc[0] + "00",
                    ddc[:2] + "0",
                    row[indexes["Title"]],
                    row[indexes["description"]],
                    row[indexes["remark"]],
                    rng.random(),
                )
            )
            counts[level] += 1
            if len(batch) >= 1000:
                connection.executemany(insert_sql, batch)
                batch.clear()
        if batch:
            connection.executemany(insert_sql, batch)
        connection.commit()
        print(f"Loaded {sheet_name}: {counts[level]:,} rows")

    workbook.close()
    connection.execute(
        "CREATE INDEX row_lookup ON rows(level,ddc,random_key)"
    )
    connection.execute(
        "CREATE INDEX dl1_child_lookup ON rows(level,parent_dl1,random_key)"
    )
    connection.execute(
        "CREATE INDEX dl2_child_lookup ON rows(level,parent_dl2,random_key)"
    )
    connection.execute(
        """
        CREATE TABLE eligible_dl3 AS
        SELECT ddc, COUNT(*) AS source_count
        FROM rows
        WHERE level=3
        GROUP BY ddc
        HAVING COUNT(*) >= ?
        """,
        (DL3_MINIMUM,),
    )
    connection.execute(
        "CREATE UNIQUE INDEX eligible_dl3_code ON eligible_dl3(ddc)"
    )
    connection.commit()
    return counts


def make_output_workbook():
    workbook = Workbook(write_only=True)
    sheets = {}
    for level in (1, 2, 3):
        worksheet = workbook.create_sheet(f"DL{level}")
        worksheet.sheet_properties.tabColor = SHEET_COLORS[level]
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        header_cells = []
        for header in OUTPUT_HEADERS:
            cell = WriteOnlyCell(worksheet, value=header)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=SHEET_COLORS[level])
            cell.alignment = Alignment(horizontal="center", vertical="center")
            header_cells.append(cell)
        worksheet.append(header_cells)
        sheets[level] = worksheet
    return workbook, sheets


def append_row(worksheet, target_ddc, source_level, source_row):
    source_ddc, title, description, remark = source_row
    ddc_cell = WriteOnlyCell(worksheet, value=target_ddc)
    ddc_cell.number_format = "@"
    ddc_cell.alignment = Alignment(horizontal="center", vertical="center")
    source_ddc_cell = WriteOnlyCell(worksheet, value=source_ddc)
    source_ddc_cell.number_format = "@"
    worksheet.append(
        [
            ddc_cell,
            title,
            description,
            remark,
            f"DL{source_level}",
            source_ddc_cell,
        ]
    )


def node_codes(connection, level):
    return [
        row[0]
        for row in connection.execute(
            "SELECT DISTINCT ddc FROM rows WHERE level=? ORDER BY ddc", (level,)
        )
    ]


def own_rows(connection, level, ddc, limit):
    return connection.execute(
        """
        SELECT ddc,title,description,remark
        FROM rows
        WHERE level=? AND ddc=?
        ORDER BY random_key,id
        LIMIT ?
        """,
        (level, ddc, limit),
    ).fetchall()


def dl2_child_rows(connection, dl1_ddc, limit):
    return connection.execute(
        """
        SELECT ddc,title,description,remark
        FROM rows
        WHERE level=2 AND parent_dl1=?
        ORDER BY random_key,id
        LIMIT ?
        """,
        (dl1_ddc, limit),
    ).fetchall()


def eligible_dl3_child_rows(connection, dl2_ddc, limit):
    return connection.execute(
        """
        SELECT r.ddc,r.title,r.description,r.remark
        FROM rows AS r
        JOIN eligible_dl3 AS e ON e.ddc=r.ddc
        WHERE r.level=3 AND r.parent_dl2=?
        ORDER BY r.random_key,r.id
        LIMIT ?
        """,
        (dl2_ddc, limit),
    ).fetchall()


def build_dl1(connection, worksheet):
    target = TARGETS[1]
    written = supplemented = 0
    for ddc in node_codes(connection, 1):
        own = own_rows(connection, 1, ddc, target)
        needed = target - len(own)
        children = dl2_child_rows(connection, ddc, needed) if needed else []
        if len(own) + len(children) != target:
            raise RuntimeError(
                f"DL1 {ddc}: only {len(own)} own + {len(children)} child rows; "
                f"cannot reach {target}"
            )
        for row in own:
            append_row(worksheet, ddc, 1, row)
        for row in children:
            append_row(worksheet, ddc, 2, row)
        written += target
        supplemented += len(children)
    return written, supplemented


def build_dl2(connection, worksheet):
    target = TARGETS[2]
    written = supplemented = 0
    for ddc in node_codes(connection, 2):
        own = own_rows(connection, 2, ddc, target)
        needed = target - len(own)
        children = eligible_dl3_child_rows(connection, ddc, needed) if needed else []
        if len(own) + len(children) != target:
            raise RuntimeError(
                f"DL2 {ddc}: only {len(own)} own + {len(children)} eligible "
                f"DL3 rows; cannot reach {target}"
            )
        for row in own:
            append_row(worksheet, ddc, 2, row)
        for row in children:
            append_row(worksheet, ddc, 3, row)
        written += target
        supplemented += len(children)
    return written, supplemented


def build_dl3(connection, worksheet):
    written = 0
    retained_nodes = 0
    for ddc, source_count in connection.execute(
        "SELECT ddc,source_count FROM eligible_dl3 ORDER BY ddc"
    ):
        keep = min(source_count, TARGETS[3])
        rows = own_rows(connection, 3, ddc, keep)
        for row in rows:
            append_row(worksheet, ddc, 3, row)
        written += len(rows)
        retained_nodes += 1
    source_nodes = connection.execute(
        "SELECT COUNT(DISTINCT ddc) FROM rows WHERE level=3"
    ).fetchone()[0]
    return written, retained_nodes, source_nodes - retained_nodes


def format_output_sheets(sheets, output_counts):
    widths = [10, 42, 90, 24, 12, 14]
    for level, worksheet in sheets.items():
        for column, width in enumerate(widths, 1):
            worksheet.column_dimensions[get_column_letter(column)].width = width
        worksheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(OUTPUT_HEADERS))}{output_counts[level] + 1}"
        )


def main():
    args = parse_args()
    if not args.input.is_file():
        raise FileNotFoundError(args.input)
    if args.output.exists() and not args.overwrite:
        raise FileExistsError(
            f"Output already exists: {args.output}\n"
            "Use --overwrite or choose a different --output path."
        )

    started = time.perf_counter()
    temp_database = args.output.with_suffix(args.output.suffix + ".sampling.sqlite3")
    if temp_database.exists():
        temp_database.unlink()

    connection = create_database(temp_database)
    try:
        source_counts = load_input_into_database(args.input, connection, args.seed)
        workbook, sheets = make_output_workbook()
        dl1_count, dl1_supplemented = build_dl1(connection, sheets[1])
        dl2_count, dl2_supplemented = build_dl2(connection, sheets[2])
        dl3_count, dl3_nodes, dl3_dropped_nodes = build_dl3(connection, sheets[3])
        output_counts = {1: dl1_count, 2: dl2_count, 3: dl3_count}
        format_output_sheets(sheets, output_counts)

        args.output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(args.output)
    finally:
        connection.close()
        if temp_database.exists() and not args.keep_temp:
            temp_database.unlink()

    elapsed = time.perf_counter() - started
    print("\nCompleted")
    print(f"Source rows: DL1={source_counts[1]:,}, DL2={source_counts[2]:,}, DL3={source_counts[3]:,}")
    print(f"DL1 output: {dl1_count:,} rows; supplemented from DL2: {dl1_supplemented:,}")
    print(f"DL2 output: {dl2_count:,} rows; supplemented from DL3: {dl2_supplemented:,}")
    print(
        f"DL3 output: {dl3_count:,} rows across {dl3_nodes:,} nodes; "
        f"dropped nodes below {DL3_MINIMUM}: {dl3_dropped_nodes:,}"
    )
    print(f"Random seed: {args.seed}")
    print(f"Output: {args.output.resolve()}")
    print(f"Wall time: {elapsed:.2f} seconds")


if __name__ == "__main__":
    main()
