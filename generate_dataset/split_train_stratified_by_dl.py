"""Split one DDC dataset into independent DL1, DL2 and DL3 Excel sheets.

Rules:
  DL1: original one-digit DDC, right-pad to three digits (1 -> 100)
  DL2: original two-digit DDC, right-pad to three digits (12 -> 120)
  DL3: original three-digit DDC, copy unchanged (123 -> 123)

All other columns and row values are copied without combining records between
levels. The implementation streams both input and output to limit memory use.
"""

import argparse
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEET_COLORS = {
    1: "1F4E78",
    2: "4472C4",
    3: "5B9BD5",
}


def parse_args():
    base_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        type=Path,
        default=base_dir / "train_stratified.xlsx",
        help="Source XLSX (default: train_stratified.xlsx beside this script)",
    )
    parser.add_argument(
        "--sheet",
        default="Sheet1",
        help="Source sheet name (default: Sheet1)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=base_dir / "train_stratified_DL1_DL2_DL3.xlsx",
        help="New XLSX path",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow replacing an existing output file",
    )
    return parser.parse_args()


def normalize_code(raw_value, row_number):
    if raw_value is None:
        raise ValueError(f"Row {row_number}: empty DDC")

    code = str(raw_value).strip()
    if code.endswith(".0") and code[:-2].isdigit():
        code = code[:-2]
    if not code.isdigit() or len(code) not in (1, 2, 3):
        raise ValueError(
            f"Row {row_number}: DDC must be a 1-3 digit code, got {raw_value!r}"
        )

    level = len(code)
    normalized = code.ljust(3, "0") if level in (1, 2) else code
    return level, normalized


def styled_header(worksheet, headers, color):
    cells = []
    for value in headers:
        cell = WriteOnlyCell(worksheet, value=value)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cells.append(cell)
    return cells


def output_row(worksheet, source_row, ddc_index, normalized_code):
    values = list(source_row)
    ddc_cell = WriteOnlyCell(worksheet, value=normalized_code)
    ddc_cell.number_format = "@"  # Keep 000, 010, 001 visibly three digits.
    ddc_cell.alignment = Alignment(horizontal="center", vertical="center")
    values[ddc_index] = ddc_cell
    return values


def main():
    args = parse_args()
    if not args.input.is_file():
        raise FileNotFoundError(args.input)
    if args.output.exists() and not args.overwrite:
        raise FileExistsError(
            f"Output already exists: {args.output}\n"
            "Use --overwrite or choose another --output path."
        )

    started = time.perf_counter()
    source_book = load_workbook(args.input, read_only=True, data_only=True)
    if args.sheet not in source_book.sheetnames:
        raise ValueError(
            f"Sheet {args.sheet!r} not found; available: {source_book.sheetnames}"
        )
    source_sheet = source_book[args.sheet]
    rows = source_sheet.iter_rows(values_only=True)
    headers = tuple(next(rows))
    if "DDC" not in headers:
        raise ValueError(f"DDC column not found; columns: {headers}")
    ddc_index = headers.index("DDC")

    output_book = Workbook(write_only=True)
    sheets = {}
    counts = {1: 0, 2: 0, 3: 0}
    for level in (1, 2, 3):
        sheet = output_book.create_sheet(f"DL{level}")
        sheet.sheet_properties.tabColor = SHEET_COLORS[level]
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.append(styled_header(sheet, headers, SHEET_COLORS[level]))
        sheets[level] = sheet

    for row_number, row in enumerate(rows, 2):
        level, normalized_code = normalize_code(row[ddc_index], row_number)
        sheets[level].append(
            output_row(sheets[level], row, ddc_index, normalized_code)
        )
        counts[level] += 1
        processed = sum(counts.values())
        if processed % 25000 == 0:
            print(f"Processed {processed:,} rows...")

    # Column widths follow the source schema but also work for additional columns.
    preferred_widths = {
        "DDC": 10,
        "Title": 42,
        "description": 90,
        "remark": 24,
    }
    for level, sheet in sheets.items():
        for index, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                preferred_widths.get(str(header), 18)
            )
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{counts[level] + 1}"
        )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    output_book.save(args.output)
    source_book.close()

    elapsed = time.perf_counter() - started
    print("Completed")
    print(f"DL1: {counts[1]:,} rows (one digit -> three digits)")
    print(f"DL2: {counts[2]:,} rows (two digits -> three digits)")
    print(f"DL3: {counts[3]:,} rows (copied unchanged)")
    print(f"Output: {args.output.resolve()}")
    print(f"Wall time: {elapsed:.2f} seconds")


if __name__ == "__main__":
    main()
